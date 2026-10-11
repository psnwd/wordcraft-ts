"""JSON -> .docx report engine (python-docx + matplotlib), plus a sample-JSON generator.

    render_spec(spec) -> bytes      build a document from a JSON spec (see ../REPORT_JSON.md for the schema)
    make_spec(params) -> dict       synthetic sample: a big spec whose text is computed from generated numbers
    analyze_report(bytes) -> dict   outline, pictures, sections, lists, header/footer of a finished document

The spec is plain JSON, so any system can produce it: one big document in, one big report out.
"""

from __future__ import annotations

import base64
import binascii
import io
import random
import re
from dataclasses import dataclass

import matplotlib

matplotlib.use("Agg")  # no display needed
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

BLUE, VIOLET, MUTED, RED = "#2F6FE0", "#8B5CF6", "#6B7280", "#C0392B"
PALETTE = ["#C7D2FE", BLUE, VIOLET, "#F5A524", "#10B981", "#EF4444"]
NUMERIC = re.compile(r"^[-+−]?[$€£]?[\d,]*\.?\d+%?$")
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

# limits at the trust boundary: the spec is untrusted input
MAX_BLOCKS, MAX_TABLE_ROWS, MAX_TABLE_COLS, MAX_ITEMS, MAX_TEXT, MAX_CHART_POINTS = 20_000, 50_000, 30, 5_000, 20_000, 500


class SpecError(ValueError):
    """The JSON spec is invalid; the message says where."""


# ---- validation ------------------------------------------------------------------------------------------


def _text(v, where: str, *, required=False) -> str:
    if v is None or v == "":
        if required:
            raise SpecError(f"{where}: text is required")
        return ""
    if not isinstance(v, (str, int, float)) or isinstance(v, bool):
        raise SpecError(f"{where}: expected a string")
    s = str(v)
    if len(s) > MAX_TEXT:
        raise SpecError(f"{where}: text longer than {MAX_TEXT} characters")
    return s


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)[:MAX_TEXT]


def _hex(v, where: str, default: str) -> str:
    if v is None:
        return default
    if not (isinstance(v, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", v)):
        raise SpecError(f"{where}: colours look like #RRGGBB")
    return "#" + v.lstrip("#")


def validate(spec) -> dict:
    """Check the whole spec up front so a bad block 9,000 reports 'blocks[9000]', not a stack trace."""
    if not isinstance(spec, dict):
        raise SpecError("spec must be a JSON object")
    blocks = spec.get("blocks", [])
    if not isinstance(blocks, list):
        raise SpecError("blocks must be a list")
    if len(blocks) > MAX_BLOCKS:
        raise SpecError(f"too many blocks ({len(blocks)} > {MAX_BLOCKS})")
    for i, b in enumerate(blocks):
        w = f"blocks[{i}]"
        if not isinstance(b, dict) or "type" not in b:
            raise SpecError(f'{w}: needs a "type"')
        t = b["type"]
        if t == "heading":
            if b.get("level", 1) not in (1, 2, 3):
                raise SpecError(f"{w}: heading level must be 1, 2 or 3")
            _text(b.get("text"), w, required=True)
        elif t in ("paragraph", "quote"):
            _text(b.get("text"), w, required=True)
            if b.get("align") not in (None, *ALIGN):
                raise SpecError(f"{w}: align must be one of {list(ALIGN)}")
        elif t in ("bullets", "numbered"):
            items = b.get("items")
            if not isinstance(items, list) or len(items) > MAX_ITEMS:
                raise SpecError(f"{w}: items must be a list of at most {MAX_ITEMS}")
            for j, it in enumerate(items):
                _text(it, f"{w}.items[{j}]", required=True)
        elif t == "table":
            cols, rows = b.get("columns"), b.get("rows")
            if not isinstance(cols, list) or not cols or len(cols) > MAX_TABLE_COLS:
                raise SpecError(f"{w}: columns must be a non-empty list of at most {MAX_TABLE_COLS}")
            if not isinstance(rows, list) or len(rows) > MAX_TABLE_ROWS:
                raise SpecError(f"{w}: rows must be a list of at most {MAX_TABLE_ROWS}")
            for j, r in enumerate(rows):
                if not isinstance(r, list) or len(r) > len(cols):
                    raise SpecError(f"{w}.rows[{j}]: a row is a list with at most {len(cols)} cells")
            if b.get("total") is not None and not isinstance(b["total"], list):
                raise SpecError(f"{w}: total must be a list of cells")
            if b.get("align") is not None and (not isinstance(b["align"], list) or any(a not in ("left", "right", "center") for a in b["align"])):
                raise SpecError(f"{w}: align must be a list of left/right/center")
        elif t == "chart":
            if b.get("kind") not in ("bar", "line"):
                raise SpecError(f'{w}: chart kind must be "bar" or "line"')
            labels, series = b.get("labels"), b.get("series")
            if not isinstance(labels, list) or not 1 <= len(labels) <= MAX_CHART_POINTS:
                raise SpecError(f"{w}: labels must be a list of 1..{MAX_CHART_POINTS}")
            if not isinstance(series, list) or not 1 <= len(series) <= 6:
                raise SpecError(f"{w}: series must be a list of 1..6")
            for j, s in enumerate(series):
                v = s.get("values") if isinstance(s, dict) else None
                if not isinstance(v, list) or len(v) != len(labels) or any(isinstance(x, bool) or not isinstance(x, (int, float)) for x in v):
                    raise SpecError(f"{w}.series[{j}]: values must be {len(labels)} numbers")
                _hex(s.get("color"), f"{w}.series[{j}].color", BLUE)
        elif t == "image":
            try:
                base64.b64decode(b.get("data", ""), validate=True)
            except (binascii.Error, TypeError, ValueError):
                raise SpecError(f"{w}: data must be base64 PNG/JPEG bytes") from None
        elif t == "section":
            if b.get("orientation") not in ("portrait", "landscape"):
                raise SpecError(f'{w}: orientation must be "portrait" or "landscape"')
        elif t != "pageBreak":
            raise SpecError(f'{w}: unknown block type "{t}"')
    return spec


# ---- rendering helpers -----------------------------------------------------------------------------------


def shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    pr.append(shd)


def add_field(paragraph, instr: str, placeholder: str = "1") -> None:
    """A complex field (PAGE, NUMPAGES...) with a cached result, so it shows even before fields update."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, placeholder), ("end", None)):
        run = paragraph.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def set_cell(cell, text: str, *, bold=False, right=False, center=False, fill=None, color=None, size=10) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    if right:
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if fill:
        shade(cell, fill)


def add_table(doc, b: dict) -> None:
    cols = [str(c) for c in b["columns"]]
    rows = [[_cell(c) for c in r] + [""] * (len(cols) - len(r)) for r in b["rows"]]
    align = b.get("align") or [
        "right" if rows and NUMERIC.match(rows[0][i].strip() or "x") else "left" for i in range(len(cols))
    ]
    align += ["left"] * (len(cols) - len(align))
    highlight = {str(k): _hex(v, "table.highlight", RED) for k, v in (b.get("highlight") or {}).items()}
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if b.get("headerRepeat", True):
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))  # repeat on every page
    for i, (cell, h) in enumerate(zip(table.rows[0].cells, cols)):
        set_cell(cell, h, bold=True, right=align[i] == "right", center=align[i] == "center", fill="#D9E2F3")
    zebra = b.get("zebra", True)
    negatives = b.get("colourNegatives", True)
    for n, row in enumerate(rows):
        new_row = table.add_row()  # use the returned row: table.rows[-1] would rebuild every row each time (quadratic)
        new_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # keep a row on one page
        cells = new_row.cells
        fill = "#F5F7FB" if zebra and n % 2 else None
        for i, (cell, v) in enumerate(zip(cells, row)):
            color = highlight.get(v) or (RED if negatives and v.startswith(("-", "−")) and align[i] == "right" else None)
            set_cell(cell, v, right=align[i] == "right", center=align[i] == "center", fill=fill, color=color)
    if b.get("total"):
        total = [_cell(c) for c in b["total"]] + [""] * (len(cols) - len(b["total"]))
        for i, (cell, v) in enumerate(zip(table.add_row().cells, total)):
            set_cell(cell, v, bold=True, right=align[i] == "right", center=align[i] == "center", fill="#E8EEF9")
    if b.get("caption"):
        caption(doc, b["caption"])


def caption(doc, text: str) -> None:
    p = doc.add_paragraph(str(text))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True


def add_chart(doc, b: dict) -> None:
    fig, ax = plt.subplots(figsize=(7, 3.2))
    labels = [str(x) for x in b["labels"]]
    series = b["series"]
    n = len(series)
    for k, s in enumerate(series):
        color = _hex(s.get("color"), "chart.color", PALETTE[(k + 1) % len(PALETTE)] if n > 1 else BLUE)
        name = str(s.get("name", ""))
        if b["kind"] == "bar":
            width = 0.8 / n
            xs = [i + (k - (n - 1) / 2) * width for i in range(len(labels))]
            ax.bar(xs, s["values"], width, label=name or None, color=color)
        else:
            ax.plot(labels, s["values"], marker="o" if len(labels) <= 24 else None, linewidth=2.4, label=name or None, color=color)
            if n == 1:
                ax.fill_between(labels, s["values"], alpha=0.12, color=color)
    if b["kind"] == "bar":
        ax.set_xticks(list(range(len(labels))), labels)
    elif len(labels) > 12:
        ax.set_xticks(list(range(0, len(labels), max(1, len(labels) // 10))))
    if b.get("yLabel"):
        ax.set_ylabel(str(b["yLabel"]))
    if b.get("title"):
        ax.set_title(str(b["title"]), fontsize=10)
    if any(s.get("name") for s in series):
        ax.legend(frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    out = io.BytesIO()
    fig.savefig(out, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    out.seek(0)
    doc.add_picture(out, width=Inches(float(b.get("width_in", 6))))
    if b.get("caption"):
        caption(doc, b["caption"])


def render_spec(spec: dict) -> bytes:
    spec = validate(spec)
    blocks = spec.get("blocks", [])
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    for lvl, size in ((1, 20), (2, 15), (3, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name, st.font.size = "Calibri", Pt(size)
        st.font.color.rgb = RGBColor.from_string(BLUE.lstrip("#"))
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)

    cover = bool(spec.get("cover", True)) and bool(spec.get("title"))
    sec.different_first_page_header_footer = cover  # no header/footer on the cover
    if spec.get("header"):
        hp = sec.header.paragraphs[0]
        hp.text = _text(spec["header"], "header")
        hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(9), RGBColor.from_string(MUTED.lstrip("#"))
    if spec.get("pageNumbers", True):
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run("Page ")
        add_field(fp, "PAGE")
        fp.add_run(" of ")
        add_field(fp, "NUMPAGES")

    if cover:
        t = doc.add_paragraph(_text(spec["title"], "title"), style="Title")
        t.paragraph_format.space_before = Pt(170)
        if spec.get("subtitle"):
            doc.add_paragraph(_text(spec["subtitle"], "subtitle"), style="Subtitle")
        if spec.get("meta"):
            m = doc.add_paragraph(_text(spec["meta"], "meta"))
            m.runs[0].font.color.rgb = RGBColor.from_string(MUTED.lstrip("#"))
        doc.add_page_break()
    if spec.get("toc", True):
        h1 = [b["text"] for b in blocks if b["type"] == "heading" and b.get("level", 1) == 1]
        if h1:  # a static list: page numbers would need Word to lay the document out
            doc.add_heading("Contents", level=1)
            for h in h1:
                doc.add_paragraph(str(h))
            doc.add_page_break()

    for b in blocks:
        t = b["type"]
        if t == "heading":
            doc.add_heading(str(b["text"]), level=b.get("level", 1))
        elif t == "paragraph":
            p = doc.add_paragraph()
            run = p.add_run(str(b["text"]))
            run.bold, run.italic = bool(b.get("bold")), bool(b.get("italic"))
            if b.get("align"):
                p.alignment = ALIGN[b["align"]]
        elif t == "quote":
            doc.add_paragraph(str(b["text"]), style="Quote").paragraph_format.space_before = Pt(12)
        elif t in ("bullets", "numbered"):
            for it in b["items"]:
                doc.add_paragraph(str(it), style="List Bullet" if t == "bullets" else "List Number")
        elif t == "table":
            add_table(doc, b)
        elif t == "chart":
            add_chart(doc, b)
        elif t == "image":
            doc.add_picture(io.BytesIO(base64.b64decode(b["data"])), width=Inches(float(b.get("width_in", 5))))
            if b.get("caption"):
                caption(doc, b["caption"])
        elif t == "pageBreak":
            doc.add_page_break()
        elif t == "section":
            new = doc.add_section(WD_SECTION.NEW_PAGE)
            new.different_first_page_header_footer = False
            land = b["orientation"] == "landscape"
            if land != (new.orientation == WD_ORIENT.LANDSCAPE):
                new.page_width, new.page_height = new.page_height, new.page_width
            new.orientation = WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---- sample spec -----------------------------------------------------------------------------------------

ALL_REGIONS = ["North America", "EMEA", "APAC", "Latin America", "Middle East", "Africa"]
PRODUCTS = ["Studio", "Cloud Sync", "Review Suite", "Mobile", "API"]
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
NAME_A = ["Alder", "Birch", "Cedar", "Delta", "Ember", "Fjord", "Granite", "Harbor", "Iris", "Juniper", "Kestrel", "Lumen"]
NAME_B = ["Press", "Labs", "Works", "Studio", "Group", "Systems", "Co-op", "Media", "Partners", "Holdings"]
PLANS = ["Starter", "Team", "Business", "Enterprise"]
STATUSES = ["Active", "Active", "Active", "Trial", "Churned"]


@dataclass
class Params:
    company: str = "Alder Street Press"
    year: int = 2025
    seed: int = 7
    regions: int = 4  # 2..6
    appendix_rows: int = 80  # 0..50000
    charts: bool = True

    def clean(self) -> "Params":
        self.company = str(self.company).strip()[:80] or "Company"
        self.year = max(1990, min(2100, int(self.year)))
        self.seed = int(self.seed)
        self.regions = max(2, min(len(ALL_REGIONS), int(self.regions)))
        self.appendix_rows = max(0, min(MAX_TABLE_ROWS, int(self.appendix_rows)))
        self.charts = bool(self.charts)
        return self


def pct(new: float, old: float) -> float:
    return (new - old) / old if old else 0.0


def money(v: float) -> str:
    return f"{v:,.0f}"


def make_spec(params: Params) -> dict:
    """A realistic big report as JSON. Numbers are synthetic and deterministic per seed; the text is computed."""
    p = params.clean()
    rnd = random.Random(p.seed)
    regions = ALL_REGIONS[: p.regions]
    prev, cur = {}, {}
    for r in regions:
        base, g = rnd.uniform(700, 2600), rnd.gauss(0.10, 0.08)
        prev[r] = [base * rnd.uniform(0.9, 1.1) * (1 + 0.02 * i) for i in range(4)]
        cur[r] = [v * (1 + g + rnd.gauss(0, 0.03)) for v in prev[r]]
    prod = {}
    for n in PRODUCTS:
        base = rnd.uniform(1500, 6000)
        prod[n] = (base, base * (1 + rnd.gauss(0.09, 0.12)), rnd.uniform(0.55, 0.85))
    customers = [
        {"id": f"C-{1000 + i}", "name": f"{rnd.choice(NAME_A)} {rnd.choice(NAME_B)}", "region": rnd.choice(regions), "plan": rnd.choice(PLANS),
         "seats": rnd.randint(3, 900), "mrr": round(rnd.uniform(90, 14000), 2), "status": rnd.choice(STATUSES)}
        for i in range(p.appendix_rows)
    ]

    tp, tc = {r: sum(prev[r]) for r in regions}, {r: sum(cur[r]) for r in regions}
    rev_p, rev_c = sum(tp.values()), sum(tc.values())
    growth = pct(rev_c, rev_p)
    ranked = sorted(regions, key=lambda r: pct(tc[r], tp[r]), reverse=True)
    top, weak = ranked[0], ranked[-1]
    q_tot = [sum(cur[r][i] for r in regions) for i in range(4)]
    best_q = QUARTERS[q_tot.index(max(q_tot))]
    biggest = max(regions, key=lambda r: tc[r])
    share = tc[biggest] / rev_c
    gm = sum(c * m for _, c, m in prod.values()) / sum(c for _, c, _ in prod.values())
    best_p = max(prod, key=lambda n: pct(prod[n][1], prod[n][0]))
    worst_p = min(prod, key=lambda n: pct(prod[n][1], prod[n][0]))
    active = [c for c in customers if c["status"] == "Active"]

    blocks: list[dict] = [
        {"type": "heading", "level": 1, "text": "1. Executive summary"},
        {"type": "paragraph", "text": (
            f"{p.company} closed {p.year} with revenue of ${money(rev_c)}k, {'up' if growth >= 0 else 'down'} {abs(growth):.1%} "
            f"on ${money(rev_p)}k the year before. {top} grew fastest ({pct(tc[top], tp[top]):+.1%}), while {weak} moved at "
            f"{pct(tc[weak], tp[weak]):+.1%}. The strongest quarter was {best_q}.")},
        {"type": "table", "columns": ["Key figure", "Value"], "align": ["left", "right"], "rows": [
            ["Revenue ($k)", money(rev_c)], ["Year-over-year growth", f"{growth:+.1%}"], ["Blended gross margin", f"{gm:.1%}"],
            ["Largest region", f"{biggest} ({share:.0%} of revenue)"],
            ["Customers in the ledger", f"{len(customers)} ({len(active)} active)" if customers else "n/a"]]},
        {"type": "quote", "text": f"“Growth came from {top}, not from price: volumes carried the year.”"},
        {"type": "heading", "level": 1, "text": "2. Revenue analysis"},
        {"type": "heading", "level": 2, "text": "2.1 Revenue by region and quarter"},
        {"type": "paragraph", "text": "All figures are in thousands of US dollars. Negative changes are shown in red."},
        {"type": "table", "columns": ["Region", *QUARTERS, "Total", "YoY"],
         "rows": [[r, *[money(v) for v in cur[r]], money(tc[r]), f"{pct(tc[r], tp[r]):+.1%}"] for r in regions],
         "total": ["Total", *[money(v) for v in q_tot], money(rev_c), f"{growth:+.1%}"]},
    ]
    if p.charts:
        blocks += [
            {"type": "heading", "level": 2, "text": "2.2 Charts"},
            {"type": "chart", "kind": "bar", "labels": regions, "yLabel": "$k", "caption": f"Figure 1. Revenue by region, {p.year - 1} vs {p.year}.",
             "series": [{"name": str(p.year - 1), "color": "#C7D2FE", "values": [round(tp[r]) for r in regions]},
                        {"name": str(p.year), "color": BLUE, "values": [round(tc[r]) for r in regions]}]},
            {"type": "chart", "kind": "line", "labels": QUARTERS, "yLabel": "$k", "caption": f"Figure 2. Total revenue by quarter. {best_q} was the peak.",
             "series": [{"name": "", "color": VIOLET, "values": [round(v) for v in q_tot]}]},
        ]
    blocks.append({"type": "heading", "level": 2, "text": "2.3 What moved"})
    blocks.append({"type": "bullets", "items": [f"{r}: {pct(tc[r], tp[r]):+.1%} to ${money(tc[r])}k." for r in ranked[:3]]})
    blocks.append({"type": "heading", "level": 2, "text": "2.4 Regional commentary"})
    for r in regions:
        qs = cur[r]
        vs = pct(tc[r], tp[r]) - growth
        blocks += [
            {"type": "heading", "level": 3, "text": r},
            {"type": "paragraph", "text": (
                f"{r} delivered ${money(tc[r])}k, {pct(tc[r], tp[r]):+.1%} year over year and {tc[r] / rev_c:.0%} of group revenue. "
                f"Its best quarter was {QUARTERS[qs.index(max(qs))]} (${money(max(qs))}k) and its weakest {QUARTERS[qs.index(min(qs))]} "
                f"(${money(min(qs))}k), a spread of {(max(qs) - min(qs)) / max(qs):.0%}. That is {abs(vs):.1%} points "
                f"{'ahead of' if vs >= 0 else 'behind'} the group, "
                f"{'so it is a candidate for more investment' if vs >= 0 else 'so we are reviewing the plan for next year'}.")}]
    all_new = sum(b for _, b, _ in prod.values())
    blocks += [
        {"type": "heading", "level": 1, "text": "3. Product performance"},
        {"type": "table", "columns": ["Product", f"{p.year - 1} ($k)", f"{p.year} ($k)", "Change", "Gross margin"],
         "rows": [[n, money(a), money(b), f"{pct(b, a):+.1%}", f"{m:.0%}"] for n, (a, b, m) in prod.items()],
         "total": ["All products", money(sum(a for a, _, _ in prod.values())), money(all_new),
                   f"{pct(all_new, sum(a for a, _, _ in prod.values())):+.1%}", f"{gm:.0%}"]},
        {"type": "paragraph", "text": f"{best_p} was the standout ({pct(prod[best_p][1], prod[best_p][0]):+.1%}). "
                                      f"{worst_p} needs attention ({pct(prod[worst_p][1], prod[worst_p][0]):+.1%})."},
        {"type": "heading", "level": 2, "text": "3.1 Product notes"},
    ]
    for n, (a, b, m) in prod.items():
        if b < a:
            verdict = "Revenue fell, and it is a high-margin line, so the loss hurts profit." if m >= gm else "Revenue fell on a below-average margin, so the profit impact is limited."
        else:
            verdict = "Margin is above the blended average, so growth here is welcome." if m >= gm else "Margin is below the blended average, so growth here dilutes profit."
        blocks += [{"type": "heading", "level": 3, "text": n},
                   {"type": "paragraph", "text": f"{n} earned ${money(b)}k ({pct(b, a):+.1%}) at a {m:.0%} gross margin, "
                                                 f"contributing {b / all_new:.0%} of product revenue. {verdict}"}]
    target = rev_c * (1 + max(0.04, growth * 0.9))
    blocks += [
        {"type": "heading", "level": 1, "text": "4. Risks and mitigations"},
        {"type": "paragraph", "text": "The main risks, in the order we would address them:"},
        {"type": "numbered", "items": [
            f"Concentration. {biggest} is {share:.0%} of revenue; a slowdown there moves the whole result.",
            f"Weak region. {weak} grew {pct(tc[weak], tp[weak]):+.1%}; review pricing and partners.",
            f"Product drag. {worst_p} shrank relative to the others; decide whether to invest or sunset.",
            "Churn. Churned accounts in the ledger show where renewals are slipping." if customers else "Churn. No ledger was included, so renewal risk is not measured here.",
            "Seasonality. Results lean on one quarter; smooth the pipeline to avoid a weak start to next year."]},
        {"type": "heading", "level": 1, "text": "5. Outlook"},
        {"type": "paragraph", "text": f"We plan for ${money(target)}k in {p.year + 1} ({pct(target, rev_c):+.1%}), led by {top}. "
                                      f"Gross margin is expected to hold near {gm:.0%}."},
        {"type": "paragraph", "italic": True, "text": "Note: all numbers in this document are synthetic, generated from a seed for demonstration."},
        {"type": "section", "orientation": "landscape"},
    ]
    if customers:
        blocks += [
            {"type": "heading", "level": 1, "text": "Appendix A. Customer ledger"},
            {"type": "paragraph", "text": f"{len(customers):,} accounts. The header row repeats on every page."},
            {"type": "table", "columns": ["ID", "Customer", "Region", "Plan", "Seats", "MRR ($)", "Status"],
             "align": ["left", "left", "left", "left", "right", "right", "left"], "highlight": {"Churned": RED},
             "rows": [[c["id"], c["name"], c["region"], c["plan"], c["seats"], c["mrr"], c["status"]] for c in customers],
             "total": ["", "Total", "", "", sum(c["seats"] for c in customers), round(sum(c["mrr"] for c in customers), 2), ""]},
        ]
    blocks += [
        {"type": "heading", "level": 1, "text": "Appendix B. Methodology"},
        {"type": "paragraph", "text": "Revenue is recognised monthly and reported by quarter. Growth compares like-for-like quarters. "
                                      "Gross margin is revenue-weighted across products. Figures are rounded to the nearest thousand."},
    ]
    return {"title": p.company, "subtitle": f"Annual Report {p.year}", "meta": f"Regions: {', '.join(regions)}",
            "header": f"{p.company}  |  Annual Report {p.year}", "pageNumbers": True, "cover": True, "toc": True, "blocks": blocks}


def build_report(params: Params) -> bytes:
    return render_spec(make_spec(params))


# ---- reading it back -------------------------------------------------------------------------------------


def analyze_report(data: bytes) -> dict:
    doc = Document(io.BytesIO(data))
    outline = []
    for para in doc.paragraphs:
        m = re.match(r"heading\s*(\d)", para.style.name, re.I)
        if m and para.text.strip():
            outline.append({"level": int(m.group(1)), "text": para.text.strip()})
    body = [para.text for para in doc.paragraphs]
    sections = [
        {"orientation": "landscape" if s.orientation == WD_ORIENT.LANDSCAPE else "portrait", "width_in": round(s.page_width.inches, 2)}
        for s in doc.sections
    ]
    first = doc.sections[0]
    return {
        "outline": outline,
        "words": sum(len(t.split()) for t in body),
        "paragraphs": sum(1 for t in body if t.strip()),
        "pictures": len(doc.inline_shapes),
        "table_shapes": [{"rows": len(t.rows), "cols": len(t.columns), "header": [c.text for c in t.rows[0].cells]} for t in doc.tables],
        "sections": sections,
        "header": first.header.paragraphs[0].text if first.header.paragraphs else "",
        "footer_has_page_numbers": "PAGE" in first.footer._element.xml and "NUMPAGES" in first.footer._element.xml,
        "lists": {
            "bullets": sum(1 for para in doc.paragraphs if para.style.name.lower().startswith("list bullet")),
            "numbered": sum(1 for para in doc.paragraphs if para.style.name.lower().startswith("list number")),
        },
    }
