"""Python side of the WordCraft example.

Four endpoints, stdlib HTTP server + python-docx:

  POST /api/generate   JSON  -> .docx   Build a document with a table made dynamically from the data.
  POST /api/sample     JSON  -> JSON    Parameters -> a big sample report spec (synthetic numbers).
  POST /api/report     JSON  -> .docx   A report spec of any size -> one document (schema: ../REPORT_JSON.md).
  POST /api/analyze    .docx -> JSON    Table contents, plus outline / pictures / sections / footer fields.

The browser edits the document in WordCraft; Python generates it and reads it back.
"""

import io
import json
from dataclasses import fields
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report import Params, analyze_report, make_spec, render_spec

MAX_BODY = 64 * 1024 * 1024  # 64 MB: a report spec can be big
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PARAM_NAMES = {f.name for f in fields(Params)}


def is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def shade(cell, hex_fill: str) -> None:
    """Background colour for a table cell (python-docx has no API for it)."""
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def build_docx(spec: dict) -> bytes:
    """spec = {"title": str, "intro": str, "rows": [{col: value, ...}, ...], "totals": bool}

    Columns come from the keys of the rows (first-seen order), so any shape of data becomes a table.
    Numeric columns are right-aligned and, with "totals", summed in a last row.
    """
    rows = spec.get("rows") or []
    if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
        raise ValueError('"rows" must be a list of objects')

    columns: list[str] = []
    for r in rows:
        for k in r:
            if k not in columns:
                columns.append(k)

    doc = Document()
    doc.add_heading(str(spec.get("title") or "Report"), level=1)
    if spec.get("intro"):
        doc.add_paragraph(str(spec["intro"]))

    if columns:
        numeric = {c for c in columns if any(is_number(r.get(c)) for r in rows)}
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for cell, name in zip(table.rows[0].cells, columns):
            cell.text = ""
            run = cell.paragraphs[0].add_run(name.replace("_", " ").title())
            run.bold = True
            shade(cell, "D9E2F3")
        for r in rows:
            cells = table.add_row().cells
            for cell, c in zip(cells, columns):
                v = r.get(c, "")
                cell.text = f"{v:,.2f}" if isinstance(v, float) else str(v)
                if c in numeric:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if spec.get("totals") and numeric:
            cells = table.add_row().cells
            cells[0].text = ""
            cells[0].paragraphs[0].add_run("Total").bold = True
            for cell, c in zip(cells, columns):
                if c in numeric:
                    total = sum(r[c] for r in rows if is_number(r.get(c)))
                    cell.text = ""
                    cell.paragraphs[0].add_run(f"{total:,.2f}" if isinstance(total, float) else str(total)).bold = True
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        doc.add_paragraph("(no data)")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def analyze_docx(data: bytes) -> dict:
    doc = Document(io.BytesIO(data))
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    text = [p.text for p in doc.paragraphs]
    simple = {
        "paragraphs": sum(1 for t in text if t.strip()),
        "words": sum(len(t.split()) for t in text),
        "headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") and p.text.strip()],
        "tables": tables,
    }
    return {**analyze_report(data), **simple}  # structure (outline, pictures, sections...) + table contents


class Handler(BaseHTTPRequestHandler):
    def _send(self, status: int, body: bytes, ctype: str) -> None:
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _json(self, status: int, obj) -> None:
        self._send(status, json.dumps(obj).encode(), "application/json")

    def do_POST(self) -> None:  # noqa: N802
        length = int(self.headers.get("Content-Length") or 0)
        if length > MAX_BODY:
            return self._json(413, {"error": "body too large"})
        body = self.rfile.read(length)
        try:
            if self.path == "/api/generate":
                return self._send(200, build_docx(json.loads(body or b"{}")), DOCX)
            if self.path == "/api/sample":
                raw = json.loads(body or b"{}")
                params = Params(**{k: v for k, v in raw.items() if k in PARAM_NAMES})  # ignore unknown keys
                return self._json(200, make_spec(params))
            if self.path == "/api/report":
                return self._send(200, render_spec(json.loads(body)), DOCX)
            if self.path == "/api/analyze":
                return self._json(200, analyze_docx(body))
        except Exception as e:  # bad JSON, not a docx, wrong shape: tell the caller, keep serving
            return self._json(400, {"error": f"{type(e).__name__}: {e}"})
        self._json(404, {"error": "not found"})

    def log_message(self, fmt, *args) -> None:
        print("[py]", fmt % args, flush=True)


if __name__ == "__main__":
    # Loopback only: this is a demo server with no auth.
    server = ThreadingHTTPServer(("127.0.0.1", 8000), Handler)
    print("Python API on http://127.0.0.1:8000")
    server.serve_forever()
